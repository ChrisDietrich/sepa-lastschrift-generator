#!/usr/bin/env python3
"""SEPA-Lastschrift (pain.008.001.02) aus einer Excel-/CSV-Tabelle erzeugen.

Ersetzt das VBA-Makro aus Excel2SepaXML.xlsm (Klassen clsSepaCDD /
clsSepaDebitInfo / clsCheckIBAN). Erwartet in jedem Tabellenblatt (Name ist
egal) dieselbe Spaltenreihenfolge wie im Original:

    A Name  B Betrag  C BIC  D IBAN  E Verwendungszweck
    F EndToEndId  G Mandatsreferenz  H Datum der Mandatsunterschrift

Bei einer Excel-Datei mit mehreren Tabellenblättern wird ohne --sheet jedes
sichtbare Blatt versucht: passt das Schema nicht (keine/fehlerhafte Daten),
wird das Blatt übersprungen und am Ende gemeldet, statt die Verarbeitung
abzubrechen - für jedes passende Blatt entsteht eine eigene SEPA-Datei.

Aufruf:
    python3 sepa_lastschrift.py mitglieder.xlsx --collection-date 01.10.2026

Die Vereins-/Gläubiger-Daten stehen in einer JSON-Datei (Standard:
creditor.json neben diesem Skript), siehe creditor.example.json.
"""
from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import re
import sys
import xml.etree.ElementTree as ET
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path

NS = "urn:iso:std:iso:20022:tech:xsd:pain.008.001.02"
XSI = "http://www.w3.org/2001/XMLSchema-instance"

# Gleiche Muster wie clsGlobalData.cls im VBA-Original
RESTRICTED_ID_RE = re.compile(r"^[A-Za-z0-9+?/\-:().,' ]{1,35}$")
IBAN_RE = re.compile(r"^[A-Z]{2}[0-9]{2}[A-Za-z0-9]{1,30}$")
BIC_RE = re.compile(r"^[A-Z]{6}[A-Z2-9][A-NP-Z0-9]([A-Z0-9]{3})?$")

PATH_SPECIAL_CHARS = "~'#%&*:<>?/\\{|}"


class ValidationError(Exception):
    """Eine einzelne Feldprüfung ist fehlgeschlagen."""


class SheetNotFoundError(Exception):
    """Das per --sheet angegebene Tabellenblatt existiert nicht in der Excel-Datei."""


class NoRowsError(Exception):
    """Die Eingabedatei enthält keine einzige Datenzeile."""


class RowValidationErrors(Exception):
    """Eine oder mehrere Zeilen der Eingabedatei sind ungültig.

    Trägt die einzelnen Fehlermeldungen in .errors, damit eine GUI sie z. B.
    als Liste anzeigen kann statt nur eines zusammengefassten Texts.
    """

    def __init__(self, errors: list[str]):
        super().__init__("; ".join(errors))
        self.errors = errors


class NoValidSheetsError(Exception):
    """Im Automatik-Modus (kein --sheet angegeben) enthielt kein einziges
    Tabellenblatt der Datei gültige Lastschriftdaten.

    Trägt die Gründe pro übersprungenem Blatt in .skipped (Liste aus
    (Blattname, Grund)-Tupeln), damit die GUI sie anzeigen kann.
    """

    def __init__(self, skipped: list[tuple[str, str]]):
        reasons = "; ".join(f"{name}: {reason}" for name, reason in skipped)
        super().__init__(f"Kein Tabellenblatt enthielt gültige Lastschriftdaten ({reasons})")
        self.skipped = skipped


def replace_umlaut(text: str) -> str:
    """SEPA-Zeichensatz kennt keine Umlaute - wie replaceUmlaut() im Original."""
    table = str.maketrans({
        "ä": "ae", "Ä": "Ae", "ö": "oe", "Ö": "Oe",
        "ü": "ue", "Ü": "Ue", "ß": "ss",
    })
    return text.translate(table)


def clean_path_string(text: str) -> str:
    for ch in PATH_SPECIAL_CHARS:
        text = text.replace(ch, "")
    return text


def clean_input(text: str) -> str:
    return (text or "").strip().replace(" ", "")


def iban_is_valid(iban: str) -> bool:
    """ISO 7064 Mod-97-10 Prüfziffer, wie clsCheckIBAN.isValid()."""
    iban = iban.replace(" ", "").upper()
    if len(iban) < 5:
        return False
    rearranged = iban[4:] + iban[:4]
    try:
        numeric = "".join(str(int(ch, 36)) for ch in rearranged)
    except ValueError:
        return False
    return int(numeric) % 97 == 1


def check_restricted_id(value: str, field_name: str, max_len: int = 35) -> str:
    value = replace_umlaut(value or "")
    if not RESTRICTED_ID_RE.match(value):
        if len(value) <= max_len:
            raise ValidationError(f"{field_name} enthält ungültige Zeichen: {value!r}")
        raise ValidationError(f"{field_name} ist länger als {max_len} Zeichen")
    return value


def check_iban(value: str, field_name: str = "IBAN") -> str:
    value = clean_input(value)
    if not IBAN_RE.match(value):
        raise ValidationError(f"{field_name} entspricht nicht dem erwarteten Muster: {value!r}")
    if not iban_is_valid(value):
        raise ValidationError(f"Die Prüfziffer der {field_name} ist ungültig: {value!r}")
    return value


def check_bic(value: str, field_name: str = "BIC") -> str:
    """BIC ist optional (seit 2016 bei SEPA-Inlandslastschriften nicht mehr Pflicht)."""
    value = clean_input(value)
    if not value:
        return ""
    if not BIC_RE.match(value):
        if len(value) <= 11:
            raise ValidationError(f"{field_name} enthält ungültige Zeichen: {value!r}")
        raise ValidationError(f"{field_name} ist länger als 11 Zeichen")
    return value


def check_amount(value) -> Decimal:
    try:
        amount = Decimal(str(value)).quantize(Decimal("0.01"))
    except (InvalidOperation, TypeError):
        raise ValidationError(f"Der Betrag ist ungültig: {value!r}")
    if not (Decimal("0.01") <= amount <= Decimal("999999999.99")):
        raise ValidationError("Der Betrag muss größer 0,00 sein")
    return amount


def check_name(value: str, field_name: str = "Name") -> str:
    value = replace_umlaut(value or "")
    if not value:
        raise ValidationError(f"{field_name} darf nicht leer sein")
    if len(value) > 70:
        raise ValidationError(f"{field_name} ist länger als 70 Zeichen")
    return value


def check_purpose(value: str) -> str:
    value = replace_umlaut(value or "")
    if not value:
        raise ValidationError("Der Verwendungszweck darf nicht leer sein")
    if len(value) > 140:
        raise ValidationError("Der Verwendungszweck ist länger als 140 Zeichen")
    return value


def check_date(value, field_name: str) -> dt.date:
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if not value:
        raise ValidationError(f"{field_name} darf nicht leer sein")
    # Deutsche Punktnotation zuerst (so tippen Nutzerinnen Daten typischerweise in Excel),
    # ISO-Notation als Fallback für maschinell erzeugte Dateien.
    for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(str(value).strip(), fmt).date()
        except ValueError:
            continue
    raise ValidationError(
        f"{field_name} ist kein gültiges Datum: {value!r} "
        "(erwartet wird ein echtes Excel-Datum oder Text im Format TT.MM.JJJJ)"
    )


@dataclass
class Creditor:
    """Vereins-/Gläubigerdaten, entspricht clsAccount + Formularfeldern in frmLastschrift."""
    name: str
    iban: str
    creditor_id: str
    bic: str = ""
    currency: str = "EUR"

    @classmethod
    def load(cls, path: Path) -> "Creditor":
        data = json.loads(path.read_text(encoding="utf-8"))
        return cls(
            name=data["name"],
            iban=data["iban"],
            creditor_id=data["creditor_id"],
            bic=data.get("bic", ""),
            currency=data.get("currency", "EUR"),
        )

    def validate(self) -> "Creditor":
        name = check_name(self.name, "Der Auftraggebername")
        iban = check_iban(self.iban, "Gläubiger-IBAN")
        bic = check_bic(self.bic, "Gläubiger-BIC")
        creditor_id = check_restricted_id(self.creditor_id, "Die Gläubiger-ID")
        if len(creditor_id) > 35:
            raise ValidationError("Die Gläubiger-ID darf maximal 35 Zeichen lang sein.")
        return Creditor(name=name, iban=iban, creditor_id=creditor_id, bic=bic, currency=self.currency)


@dataclass
class DebitRow:
    row_number: int
    name: str
    amount: Decimal
    iban: str
    purpose: str
    mandate_id: str
    date_of_signature: dt.date
    bic: str = ""
    end_to_end_id: str = ""

    @classmethod
    def from_raw(cls, row_number: int, name, amount, bic, iban, purpose,
                 end_to_end_id, mandate_id, date_of_signature) -> "DebitRow":
        errors: list[str] = []

        def collect(fn, *args):
            try:
                return fn(*args)
            except ValidationError as exc:
                errors.append(str(exc))
                return None

        name_v = collect(check_name, name, "Der Name des Zahlungspflichtigen")
        amount_v = collect(check_amount, amount)
        bic_v = collect(check_bic, bic)
        iban_v = collect(check_iban, iban)
        purpose_v = collect(check_purpose, purpose)
        # EndToEndId ist optional - "NOTPROVIDED" wie im Original, wenn leer
        if end_to_end_id:
            e2e_v = collect(check_restricted_id, str(end_to_end_id), "Die EndToEndId")
        else:
            e2e_v = "NOTPROVIDED"
        if not mandate_id:
            errors.append("Die Mandatsreferenz ist leer, diese muss bei Lastschriften gefüllt sein")
            mandate_v = None
        else:
            mandate_v = collect(check_restricted_id, str(mandate_id), "Die Mandatsreferenz")
        date_v = collect(check_date, date_of_signature, "Das Datum der Unterschrift")

        if errors:
            raise ValidationError("; ".join(errors))

        return cls(
            row_number=row_number, name=name_v, amount=amount_v, bic=bic_v or "",
            iban=iban_v, purpose=purpose_v, end_to_end_id=e2e_v,
            mandate_id=mandate_v, date_of_signature=date_v,
        )


def open_workbook(path: Path):
    import openpyxl

    return openpyxl.load_workbook(path, data_only=True)


def visible_sheet_names(wb) -> list[str]:
    """Nur sichtbare Blätter - versteckte Hilfsblätter (z. B. Nachschlagelisten aus
    dem Original-Template) sollen im Automatik-Modus nicht als Datenblätter versucht
    oder als übersprungen gemeldet werden."""
    return [ws.title for ws in wb.worksheets if ws.sheet_state == "visible"]


def extract_rows(ws) -> list[tuple]:
    rows = []
    # Wie im Original: ab Zeile 2, bis Spalte A (Name) leer ist
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=1).value
        if name in (None, ""):
            break
        rows.append(tuple(ws.cell(row=r, column=c).value for c in range(1, 9)))
    return rows


def read_rows_from_csv(path: Path) -> list[tuple]:
    rows = []
    with path.open(newline="", encoding="utf-8-sig") as f:
        reader = csv.reader(f, delimiter=";")
        next(reader, None)  # Kopfzeile überspringen
        for line in reader:
            if not line or not line[0]:
                break
            line = (line + [""] * 8)[:8]
            rows.append(tuple(line))
    return rows


def rows_from_raw(raw_rows: list[tuple]) -> list[DebitRow]:
    """Validiert Rohzeilen (aus Excel oder CSV) zu DebitRow-Objekten.

    Wirft RowValidationErrors, wenn einzelne Zeilen fehlerhaft sind. Eine leere
    Eingabe ergibt bewusst eine leere Liste statt eines Fehlers - der Aufrufer
    entscheidet, ob das ein Fehler ("Datei/Blatt hat keine Daten") oder im
    Automatik-Modus nur ein Grund zum Überspringen dieses Blatts ist.
    """
    rows: list[DebitRow] = []
    errors: list[str] = []
    for i, raw in enumerate(raw_rows, start=2):
        name, amount, bic, iban, purpose, end_to_end_id, mandate_id, date_of_signature = raw
        try:
            rows.append(DebitRow.from_raw(
                i, name, amount, bic, iban, purpose, end_to_end_id, mandate_id, date_of_signature,
            ))
        except ValidationError as exc:
            errors.append(f"Zeile {i}: {exc}")

    if errors:
        raise RowValidationErrors(errors)
    return rows


def build_pain008(creditor: Creditor, rows: list[DebitRow], *, message_id: str,
                   payment_id: str, collection_date: dt.date, sequence_type: str,
                   instrument_code: str, batch_booking: bool) -> ET.ElementTree:
    ET.register_namespace("", NS)
    ET.register_namespace("xsi", XSI)

    def se(parent, tag, text=None):
        el = ET.SubElement(parent, f"{{{NS}}}{tag}")
        if text is not None:
            el.text = text
        return el

    root = ET.Element(f"{{{NS}}}Document")
    root.set(f"{{{XSI}}}schemaLocation", f"{NS} pain.008.001.02.xsd")
    initn = se(root, "CstmrDrctDbtInitn")

    total = sum((r.amount for r in rows), Decimal("0.00"))

    grp_hdr = se(initn, "GrpHdr")
    se(grp_hdr, "MsgId", message_id)
    se(grp_hdr, "CreDtTm", dt.datetime.now().strftime("%Y-%m-%dT%H:%M:%S"))
    se(grp_hdr, "NbOfTxs", str(len(rows)))
    se(grp_hdr, "CtrlSum", str(total))
    initg_pty = se(grp_hdr, "InitgPty")
    se(initg_pty, "Nm", creditor.name)

    pmt_inf = se(initn, "PmtInf")
    se(pmt_inf, "PmtInfId", payment_id)
    se(pmt_inf, "PmtMtd", "DD")
    se(pmt_inf, "BtchBookg", "true" if batch_booking else "false")
    se(pmt_inf, "NbOfTxs", str(len(rows)))
    se(pmt_inf, "CtrlSum", str(total))

    pmt_tp_inf = se(pmt_inf, "PmtTpInf")
    svc_lvl = se(pmt_tp_inf, "SvcLvl")
    se(svc_lvl, "Cd", "SEPA")
    lcl_instrm = se(pmt_tp_inf, "LclInstrm")
    se(lcl_instrm, "Cd", instrument_code)
    se(pmt_tp_inf, "SeqTp", sequence_type)

    se(pmt_inf, "ReqdColltnDt", collection_date.strftime("%Y-%m-%d"))

    cdtr = se(pmt_inf, "Cdtr")
    se(cdtr, "Nm", creditor.name)

    cdtr_acct = se(pmt_inf, "CdtrAcct")
    cdtr_acct_id = se(cdtr_acct, "Id")
    se(cdtr_acct_id, "IBAN", creditor.iban)

    cdtr_agt = se(pmt_inf, "CdtrAgt")
    cdtr_agt_fin = se(cdtr_agt, "FinInstnId")
    if creditor.bic:
        se(cdtr_agt_fin, "BIC", creditor.bic)
    else:
        othr = se(cdtr_agt_fin, "Othr")
        se(othr, "Id", "NOTPROVIDED")

    se(pmt_inf, "ChrgBr", "SLEV")

    cdtr_schme_id = se(pmt_inf, "CdtrSchmeId")
    schme_id_id = se(cdtr_schme_id, "Id")
    prvt_id = se(schme_id_id, "PrvtId")
    othr = se(prvt_id, "Othr")
    se(othr, "Id", creditor.creditor_id)
    schme_nm = se(othr, "SchmeNm")
    se(schme_nm, "Prtry", "SEPA")

    for row in rows:
        tx = se(pmt_inf, "DrctDbtTxInf")
        pmt_id = se(tx, "PmtId")
        se(pmt_id, "EndToEndId", row.end_to_end_id)

        instd_amt = se(tx, "InstdAmt", str(row.amount))
        instd_amt.set("Ccy", creditor.currency)

        drct_dbt_tx = se(tx, "DrctDbtTx")
        mndt = se(drct_dbt_tx, "MndtRltdInf")
        se(mndt, "MndtId", row.mandate_id)
        se(mndt, "DtOfSgntr", row.date_of_signature.strftime("%Y-%m-%d"))

        dbtr_agt = se(tx, "DbtrAgt")
        dbtr_agt_fin = se(dbtr_agt, "FinInstnId")
        if row.bic:
            se(dbtr_agt_fin, "BIC", row.bic)
        else:
            othr = se(dbtr_agt_fin, "Othr")
            se(othr, "Id", "NOTPROVIDED")

        dbtr = se(tx, "Dbtr")
        se(dbtr, "Nm", row.name)

        dbtr_acct = se(tx, "DbtrAcct")
        dbtr_acct_id = se(dbtr_acct, "Id")
        se(dbtr_acct_id, "IBAN", row.iban)

        rmt_inf = se(tx, "RmtInf")
        se(rmt_inf, "Ustrd", row.purpose)

    tree = ET.ElementTree(root)
    ET.indent(tree, space="  ")
    return tree


def default_id() -> str:
    return dt.datetime.now().strftime("%Y%m%d-%H%M%S")


@dataclass
class GenerationResult:
    output_path: Path
    row_count: int
    total_amount: Decimal
    currency: str
    sheet_name: str | None = None


@dataclass
class GenerationSummary:
    results: list[GenerationResult]
    # (Blattname, Grund) für Blätter, die im Automatik-Modus übersprungen wurden,
    # weil ihr Inhalt nicht wie eine Lastschrift-Tabelle aussah.
    skipped: list[tuple[str, str]]


def generate_sepa_file(
    *,
    input_path: Path,
    creditor_path: Path,
    collection_date: dt.date,
    sequence_type: str = "RCUR",
    instrument_code: str = "CORE",
    batch_booking: bool = False,
    message_id: str | None = None,
    payment_id: str | None = None,
    out_dir: Path | None = None,
    sheet_name: str | None = None,
) -> GenerationSummary:
    """Gemeinsamer Ablauf für CLI (main()) und GUI (gui.py): Gläubigerdaten laden,
    Zeilen einlesen und prüfen, XML bauen und schreiben.

    Bei einer CSV-Datei oder explizit angegebenem `sheet_name` wird genau eine
    SEPA-Datei erzeugt; Fehler (leer/ungültig) brechen die Verarbeitung ab.

    Ohne `sheet_name` (Automatik-Modus, nur bei Excel-Dateien) wird jedes
    sichtbare Tabellenblatt versucht: enthält es gültige Lastschriftdaten, wird
    eine eigene SEPA-Datei dafür erzeugt; enthält es gar keine oder fehlerhafte
    Daten (falsches Schema), wird es übersprungen und der Grund in
    `GenerationSummary.skipped` vermerkt, ohne die übrigen Blätter zu blockieren.
    Enthält am Ende kein einziges Blatt gültige Daten, wird NoValidSheetsError
    geworfen.

    Wirft außerdem FileNotFoundError/KeyError/ValidationError (Gläubigerdatei),
    SheetNotFoundError (explizit angegebenes Blatt existiert nicht),
    RowValidationErrors oder NoRowsError (CSV oder explizites Blatt ohne Daten) -
    der Aufrufer entscheidet, wie er das anzeigt (stderr bei der CLI,
    Dialogfenster bei der GUI).
    """
    creditor = Creditor.load(creditor_path).validate()
    out_dir = out_dir or input_path.parent
    out_dir.mkdir(parents=True, exist_ok=True)

    def write_file(rows: list[DebitRow], label: str | None, filename: str, suffix: str | None) -> GenerationResult:
        mid = message_id or default_id()
        pid = payment_id or default_id()
        if suffix:
            mid, pid = f"{mid}-{suffix}", f"{pid}-{suffix}"
        mid = check_restricted_id(mid, "Die Message-ID")
        pid = check_restricted_id(pid, "Die Payment-Information-ID")

        tree = build_pain008(
            creditor, rows,
            message_id=mid, payment_id=pid, collection_date=collection_date,
            sequence_type=sequence_type, instrument_code=instrument_code,
            batch_booking=batch_booking,
        )
        out_path = out_dir / filename
        tree.write(out_path, encoding="UTF-8", xml_declaration=True)
        total = sum((r.amount for r in rows), Decimal("0.00"))
        return GenerationResult(
            output_path=out_path, row_count=len(rows), total_amount=total,
            currency=creditor.currency, sheet_name=label,
        )

    if input_path.suffix.lower() == ".csv":
        rows = rows_from_raw(read_rows_from_csv(input_path))
        if not rows:
            raise NoRowsError("Keine Einzelinformationen gefunden - es wurde keine SEPA-Datei erstellt.")
        result = write_file(rows, None, f"{input_path.stem}SEPA.xml", suffix=None)
        return GenerationSummary(results=[result], skipped=[])

    wb = open_workbook(input_path)

    if sheet_name is not None:
        if sheet_name not in wb.sheetnames:
            raise SheetNotFoundError(
                f"Blatt '{sheet_name}' nicht gefunden. Vorhandene Blätter: {', '.join(wb.sheetnames)}"
            )
        rows = rows_from_raw(extract_rows(wb[sheet_name]))
        if not rows:
            raise NoRowsError("Keine Einzelinformationen gefunden - es wurde keine SEPA-Datei erstellt.")
        result = write_file(rows, sheet_name, f"{input_path.stem}SEPA.xml", suffix=None)
        return GenerationSummary(results=[result], skipped=[])

    # Automatik-Modus: jedes sichtbare Blatt versuchen, unpassende überspringen
    # statt abzubrechen - siehe Docstring oben.
    sheets = visible_sheet_names(wb)
    multiple = len(sheets) > 1
    results: list[GenerationResult] = []
    skipped: list[tuple[str, str]] = []

    for index, name in enumerate(sheets, start=1):
        try:
            rows = rows_from_raw(extract_rows(wb[name]))
        except RowValidationErrors as exc:
            skipped.append((name, f"fehlerhafte Daten ({'; '.join(exc.errors)})"))
            continue
        if not rows:
            skipped.append((name, "keine Lastschriftdaten gefunden (Schema passt nicht)"))
            continue

        filename = (
            f"{input_path.stem}SEPA.xml" if not multiple
            else f"{input_path.stem}SEPA_{clean_path_string(name)}.xml"
        )
        suffix = str(index) if multiple else None
        results.append(write_file(rows, name, filename, suffix=suffix))

    if not results:
        raise NoValidSheetsError(skipped)

    return GenerationSummary(results=results, skipped=skipped)


class _HelpFormatter(argparse.RawDescriptionHelpFormatter, argparse.ArgumentDefaultsHelpFormatter):
    """Zeigt den mehrzeiligen Modul-Docstring unverändert an und hängt bei jeder
    Option automatisch '(Standard: ...)' an, sofern ein Default existiert."""


EPILOG = """\
Erläuterung --sequence-type (SeqTp im SEPA-XML):
  FRST  Erstlastschrift eines wiederkehrenden Mandats (z. B. erster Kursgebühren-Einzug
        eines neuen Mitglieds). Seit November 2016 ist diese Unterscheidung fachlich
        nicht mehr zwingend nötig - die Deutsche Kreditwirtschaft empfiehlt, auch für
        den ersten Einzug bereits RCUR zu verwenden.
  RCUR  Folgelastschrift eines wiederkehrenden Mandats. Praktischer Normalfall für
        laufende Kursgebühren/Mitgliedsbeiträge - deshalb der Standardwert hier.
  OOFF  Einmalige Lastschrift, für die kein wiederkehrendes Mandat vorliegt (z. B.
        ein einzelner Sonderbeitrag).
  FNAL  Letzte Lastschrift zu einem Mandat, das danach nicht mehr genutzt wird.

Erläuterung --instrument (LclInstrm/Cd im SEPA-XML):
  CORE  SEPA-Basislastschrift. Normalfall für private Mitglieder/Zahlungspflichtige.
  B2B   SEPA-Firmenlastschrift. Nur zulässig, wenn der Zahlungspflichtige seiner Bank
        ausdrücklich bestätigt hat, B2B-Mandate zu akzeptieren (typischerweise nur bei
        Firmenkonten) - für Vereinsmitglieder als Privatpersonen in aller Regel falsch.

Erläuterung --collection-date (ReqdColltnDt im SEPA-XML):
  Der Tag, an dem der Betrag vom Konto der Zahlungspflichtigen abgebucht werden soll
  ("Fälligkeitstag"). Die Datei muss vor diesem Datum bei der Bank eingereicht werden -
  wie viele Tage vorher (Vorlagefrist) hängt vom Institut und vom Instrument ab; bei der
  Sparkasse im Online-Banking nachschauen bzw. dort erfragen. Übliche Faustregeln sind
  mindestens 1 Bankarbeitstag (COR1/Folgelastschrift) bzw. mehrere Tage bei Erstlast-
  schriften - im Zweifel lieber ein paar Tage Puffer einplanen.

Beispiel:
  python3 sepa_lastschrift.py mitglieder.xlsx --collection-date 01.10.2026
"""


def parse_args(argv: list[str]) -> argparse.Namespace:
    p = argparse.ArgumentParser(
        description=__doc__,
        epilog=EPILOG,
        formatter_class=_HelpFormatter,
    )
    p.add_argument(
        "input", type=Path,
        help="Excel- (.xlsx/.xlsm) oder CSV-Datei mit den Lastschriftzeilen (Spalten: Name, "
             "Betrag, BIC, IBAN, Verwendungszweck, EndToEndId, Mandatsreferenz, Datum "
             "Mandatsunterschrift - der Tabellenblattname ist beliebig, siehe --sheet)",
    )
    p.add_argument(
        "--sheet", default=None, metavar="BLATT",
        help="Name eines bestimmten Tabellenblatts (nur bei Excel-Dateien relevant, wird bei .csv "
             "ignoriert). Ohne Angabe wird automatisch jedes sichtbare Tabellenblatt verarbeitet - "
             "je eine SEPA-Datei pro Blatt mit gültigen Lastschriftdaten; Blätter ohne passende "
             "Daten (falsches Schema) werden übersprungen und am Ende gemeldet",
    )
    p.add_argument(
        "--creditor", type=Path, default=Path(__file__).with_name("creditor.json"),
        help="JSON-Datei mit den Vereins-/Gläubigerdaten (Name, IBAN, Gläubiger-ID, optional BIC/"
             "Währung) - siehe creditor.example.json als Vorlage zum Kopieren",
    )
    p.add_argument(
        "--collection-date", required=True, metavar="DATUM",
        help="Fälligkeitstag / Einzugsdatum, an dem der Betrag abgebucht werden soll. Format "
             "TT.MM.JJJJ (z. B. 01.10.2026) oder JJJJ-MM-TT. Ausführliche Erklärung siehe unten. "
             "Pflichtangabe, kein Standardwert.",
    )
    p.add_argument(
        "--sequence-type", default="RCUR", choices=["FRST", "RCUR", "OOFF", "FNAL"],
        help="Art der Lastschrift innerhalb eines wiederkehrenden Mandats. Ausführliche "
             "Erklärung der vier Werte siehe unten",
    )
    p.add_argument(
        "--instrument", default="CORE", choices=["CORE", "B2B"],
        help="SEPA-Lastschriftverfahren. Ausführliche Erklärung siehe unten",
    )
    p.add_argument(
        "--batch-booking", action="store_true",
        help="Alle Buchungen als eine Sammelbuchung zusammenfassen. Ohne diese Option wird jede "
             "Buchung einzeln auf dem Kontoauszug ausgewiesen (Einzelbuchung, Standard)",
    )
    p.add_argument(
        "--message-id", default=None, metavar="ID",
        help="Eindeutige Kennung der Nachricht (MsgId im SEPA-XML), max. 35 Zeichen. Ohne Angabe "
             "wird automatisch ein Zeitstempel verwendet (JJJJMMTT-HHMMSS), wie im Original-Makro",
    )
    p.add_argument(
        "--payment-id", default=None, metavar="ID",
        help="Eindeutige Kennung des Zahlungsauftrags (PmtInfId im SEPA-XML), max. 35 Zeichen. "
             "Ohne Angabe wird automatisch ein Zeitstempel verwendet (JJJJMMTT-HHMMSS)",
    )
    p.add_argument(
        "--out", type=Path, default=None, metavar="ORDNER",
        help="Ordner, in dem die SEPA-XML-Datei abgelegt wird. Ohne Angabe wird derselbe Ordner "
             "wie die Eingabedatei verwendet",
    )
    return p.parse_args(argv)


def main(argv: list[str] | None = None) -> int:
    args = parse_args(argv if argv is not None else sys.argv[1:])

    try:
        collection_date = check_date(args.collection_date, "--collection-date")
    except ValidationError as exc:
        print(str(exc), file=sys.stderr)
        return 1

    try:
        summary = generate_sepa_file(
            input_path=args.input,
            creditor_path=args.creditor,
            collection_date=collection_date,
            sequence_type=args.sequence_type,
            instrument_code=args.instrument,
            batch_booking=args.batch_booking,
            message_id=args.message_id,
            payment_id=args.payment_id,
            out_dir=args.out,
            sheet_name=args.sheet,
        )
    except FileNotFoundError:
        print(f"Gläubiger-Datei nicht gefunden: {args.creditor}", file=sys.stderr)
        print("Siehe creditor.example.json als Vorlage.", file=sys.stderr)
        return 1
    except (KeyError, ValidationError) as exc:
        print(f"Fehler in der Gläubiger-Datei {args.creditor}: {exc}", file=sys.stderr)
        return 1
    except SheetNotFoundError as exc:
        print(str(exc), file=sys.stderr)
        return 1
    except RowValidationErrors as exc:
        print("Es sind Fehler in den Daten aufgetreten - es wurde keine SEPA-Datei erstellt:", file=sys.stderr)
        for e in exc.errors:
            print(f"  - {e}", file=sys.stderr)
        return 1
    except NoRowsError as exc:
        print(str(exc), file=sys.stderr)
        return 1
    except NoValidSheetsError as exc:
        print("Kein Tabellenblatt enthielt gültige Lastschriftdaten:", file=sys.stderr)
        for name, reason in exc.skipped:
            print(f"  - {name}: {reason}", file=sys.stderr)
        return 1

    for result in summary.results:
        label = f" (Blatt '{result.sheet_name}')" if result.sheet_name else ""
        print(f"SEPA-XML-Datei erstellt{label}: {result.output_path}")
        print(f"  Anzahl Datensätze: {result.row_count}")
        print(f"  Summe Datensätze: {result.total_amount:.2f} {result.currency}")

    if summary.skipped:
        print("\nÜbersprungene Tabellenblätter (kein passendes Schema):")
        for name, reason in summary.skipped:
            print(f"  - {name}: {reason}")

    return 0


if __name__ == "__main__":
    raise SystemExit(main())
